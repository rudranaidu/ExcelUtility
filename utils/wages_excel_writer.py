from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def write_wages_report(df, output_buffer, month_label):
    wb = Workbook()
    ws = wb.active

    bold = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )

    # ---------------- Title ----------------
    ws.merge_cells("A1:W1")
    ws["A1"] = "FORM-XVII REGISTER OF WAGES"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = center

    ws.merge_cells("A2:W2")
    ws["A2"] = "VIDE RULE 78 (I) A (i) OF CONTRACT LABOUR (REG. & ABOLITION) CENTRAL & A.P. RULES"
    ws["A2"].font = bold
    ws["A2"].alignment = center

    ws.merge_cells("A4:M4")
    ws["A4"] = "Name and Address of Contractor ::  PAVANII ENTERPRISES"
    ws["A4"].font = bold

    ws.merge_cells("N4:W4")
    ws["N4"] = "Name and Address of Establishment in / under  ::  GEMINI EDIBLES AND FATS LTD."
    ws["N4"].font = bold

    ws.merge_cells("A5:M5")
    ws["A5"] = "Nature and Location of Work  ::  EPURU-1A, MUTHUKURU, SPSR NELLORE, PH:9550999951."
    ws["A5"].font = bold

    ws.merge_cells("N5:W5")
    ws["N5"] = "which contract is carried on  ______________________________"
    ws["N5"].font = bold

    ws.merge_cells("A6:M6")
    ws["A6"] = f"Wage-period : Monthly  ::  {month_label}"
    ws["A6"].font = bold

    ws.merge_cells("N6:W6")
    ws["N6"] = "Name and Address of Principal Employer  ::  EPURU-1A, MUTHUKURU, SPSR NELLORE."
    ws["N6"].font = bold

    # ---------------- Table Header ----------------
    headers = [
        "Sl. No.", "Name of the workman", "Emp", "UAN NUMBER", "ESI NUMBER",
        "Bank Acc. No", "Bank Name", "Designation / Nature of work done",
        "No. of days worked", "Rate per day", "Basic Wages",
        "Leave with wages (LWW)", "Bonus @8.33%", "National Festival Holidays (NFH)",
        "Washing Allowances", "Gross", "ESI", "PF", "ESI", "Transportation",
        "PT", "TAKE HOME", "Signature"
    ]

    start_row = 9

    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=c, value=h)
        cell.font = bold
        cell.alignment = center
        cell.border = border

    # ---------------- Data Rows ----------------
    for i, row in df.iterrows():
        r = start_row + 1 + i

        def v(*keys):
            for k in keys:
                if k in row and str(row[k]) != "nan":
                    return row[k]
            return ""

        ordered = [
            i + 1,
            v("Name"),
            v("EmpID"),
            v("UAN"),
            v("ESI_NUMBER"),
            v("BankAccount"),
            v("BankName"),
            v("Category"),
            v("NoOfDays"),
            v("RatePerDay"),
            v("Basic"),
            v("LWW"),
            v("Bonus"),
            v("NFH"),
            v("Washing"),
            v("Gross"),
            v("ESI"),
            v("PF"),
            v("ESI"),
            v("Transport"),
            v("PT"),
            v("TakeHome"),
            ""
        ]

        for c, value in enumerate(ordered, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.alignment = center
            cell.border = border

    # ---------------- Column Widths ----------------
    widths = [6, 20, 10, 14, 12, 14, 14, 18, 12, 12, 12, 14, 12, 16, 14, 10, 10, 10, 10, 14, 8, 12, 12]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    wb.save(output_buffer)

