from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, Border, Side
from openpyxl.utils import get_column_letter


def write_muster_roll(df, output_buffer, month):
    wb = Workbook()
    ws = wb.active

    bold = Font(bold=True)
    big_bold = Font(bold=True, size=14)
    center = Alignment(horizontal="center", vertical="center")
    border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    total_cols = len(df.columns)
    end_col_letter = get_column_letter(total_cols)

    # ---------------- TITLE BLOCK ----------------
    ws.merge_cells(f"A1:{end_col_letter}1")
    ws["A1"] = "FORM-XVI - MUSTER ROLL"
    ws["A1"].font = big_bold
    ws["A1"].alignment = center

    ws.merge_cells(f"A2:{end_col_letter}2")
    ws["A2"] = "VIDE RULE 78 (I) A (i) OF CONTRACT LABOUR (REG. & ABOLITION) CENTRAL / A.P. RULES"
    ws["A2"].font = bold
    ws["A2"].alignment = center

    ws.merge_cells("A4:O4")
    ws["A4"] = "Name and Address of Contractor ::  PAVANII ENTERPRISES"
    ws["A4"].font = bold

    ws.merge_cells(f"P4:{end_col_letter}4")
    ws["P4"] = "Name and Address of the Establishment in/under which contract is carries on   ::  GEMINI EDIBLES AND FATS LTD."
    ws["P4"].font = bold

    ws.merge_cells("A5:O5")
    ws["A5"] = "Nature and Location od Work   ::  EPURU-1A, MUTHUKURU, SPSR NELLORE, PH:9550999951."
    ws["A5"].font = bold

    ws.merge_cells(f"P5:{end_col_letter}5")
    ws["P5"] = "Name and Address of Principal Employer  ::  EPURU-1A, MUTHUKURU, SPSR NELLORE."
    ws["P5"].font = bold

    ws.merge_cells(f"P6:{end_col_letter}6")
    ws["P6"] = f"For the Month of  ::  {month}"
    ws["P6"].font = bold

    start_row = 8

    # ---------------- COLUMN WIDTHS ----------------
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 30

    for i in range(3, total_cols):
        ws.column_dimensions[get_column_letter(i)].width = 4

    ws.column_dimensions[get_column_letter(total_cols)].width = 10

    # ---------------- TABLE HEADER ----------------
    for c, header in enumerate(df.columns, start=1):
        cell = ws.cell(row=start_row, column=c, value=str(header))
        cell.font = bold
        cell.alignment = center
        cell.border = border

    # ---------------- DATA ROWS ----------------
    for r, row in enumerate(df.itertuples(index=False), start=start_row + 1):
        for c, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.font = bold     # EVERYTHING bold
            cell.alignment = center
            cell.border = border

    # ---------------- GRAND TOTAL ----------------
    present_col = total_cols
    last_data_row = start_row + len(df)

    grand_total = df.iloc[:, -1].sum()

    total_cell = ws.cell(row=last_data_row + 1, column=present_col, value=grand_total)
    total_cell.font = bold
    total_cell.alignment = center
    total_cell.border = border

    wb.save(output_buffer)

